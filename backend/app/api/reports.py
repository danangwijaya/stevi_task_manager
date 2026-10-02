import io
import csv
from datetime import datetime
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.db.session import get_db
from app.db.models import User, TaskGrid, StudyArea, TaskStatus
from app.api.deps import get_current_user, get_current_active_reviewer
from pydantic import BaseModel

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    openpyxl = None
    HAS_OPENPYXL = False

router = APIRouter()

class TaskStageUpdateRequest(BaseModel):
    qc1_approved: Optional[bool] = None
    qc2_approved: Optional[bool] = None
    finishing_approved: Optional[bool] = None
    reviewer_notes: Optional[str] = None
    status: Optional[str] = None

def parse_grid_number(grid_code: str) -> str:
    """Extract clean grid number from grid code, e.g. SB_GRID_049_2025 -> 49"""
    if not grid_code:
        return ""
    parts = grid_code.split("_")
    for p in parts:
        if p.isdigit() and len(p) <= 4:
            return str(int(p))
    return grid_code

@router.get("/progress-table")
def get_progress_table(
    study_area_id: Optional[int] = None,
    year: Optional[int] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Mengambil data rekap tabel monitoring progres per annotator / mapper
    persis seperti format Google Spreadsheet KemenLH / UGM.
    """
    study_areas = db.query(StudyArea).order_by(StudyArea.id.asc()).all()
    study_area_map = {sa.id: sa.name for sa in study_areas}
    
    distinct_years = [y[0] for y in db.query(TaskGrid.year).distinct().order_by(TaskGrid.year.desc()).all() if y[0]]
    if not distinct_years:
        distinct_years = [2025, 2026]

    users = db.query(User).filter(User.is_active == True).order_by(User.full_name.asc()).all()

    task_query = db.query(TaskGrid).filter(TaskGrid.assigned_user_id.isnot(None))
    if study_area_id is not None and isinstance(study_area_id, int):
        task_query = task_query.filter(TaskGrid.study_area_id == study_area_id)
    if year is not None and isinstance(year, int):
        task_query = task_query.filter(TaskGrid.year == year)
    
    tasks = task_query.order_by(TaskGrid.study_area_id.asc(), TaskGrid.grid_code.asc()).all()

    user_tasks_map = {}
    for t in tasks:
        uid = t.assigned_user_id
        if uid not in user_tasks_map:
            user_tasks_map[uid] = []
        user_tasks_map[uid].append(t)

    records = []
    total_mappers_count = 0
    total_grids_count = 0
    in_progress_count = 0
    review_qc_count = 0
    qc1_count = 0
    qc2_count = 0
    finishing_count = 0

    search_lower = search.strip().lower() if (search and isinstance(search, str)) else ""

    for u in users:
        u_tasks = user_tasks_map.get(u.id, [])
        if not u_tasks and search_lower:
            continue

        grid_items = []
        user_matches_search = (
            search_lower in (u.full_name or "").lower() or
            search_lower in (u.nim_nip or "").lower() or
            search_lower in (u.department or "").lower()
        )

        for t in u_tasks:
            grid_num = parse_grid_number(t.grid_code)
            grid_matches_search = search_lower in t.grid_code.lower() or search_lower in grid_num.lower()

            if search_lower and not user_matches_search and not grid_matches_search:
                continue

            status_str = (t.status or "").upper()
            is_in_progress = status_str in ["IN_PROGRESS", "SUBMITTED", "APPROVED"]
            is_review_qc = status_str in ["SUBMITTED", "REVISION_NEEDED", "APPROVED"]
            is_qc1 = bool(t.qc1_approved or status_str == "APPROVED")
            is_qc2 = bool(t.qc2_approved)
            is_finishing = bool(t.finishing_approved)

            if is_in_progress:
                in_progress_count += 1
            if is_review_qc:
                review_qc_count += 1
            if is_qc1:
                qc1_count += 1
            if is_qc2:
                qc2_count += 1
            if is_finishing:
                finishing_count += 1

            grid_items.append({
                "task_id": t.id,
                "grid_code": t.grid_code,
                "grid_number": grid_num,
                "year": t.year,
                "study_area_id": t.study_area_id,
                "study_area_name": study_area_map.get(t.study_area_id, f"Area #{t.study_area_id}"),
                "status": t.status,
                "in_progress": is_in_progress,
                "review_qc": is_review_qc,
                "qc1_approved": is_qc1,
                "qc2_approved": is_qc2,
                "finishing_approved": is_finishing,
                "keterangan": t.reviewer_notes or "",
                "updated_at": t.updated_at.isoformat() if t.updated_at else None
            })

        if grid_items:
            total_mappers_count += 1
            total_grids_count += len(grid_items)
            records.append({
                "user_id": u.id,
                "nama": u.full_name,
                "nim": u.nim_nip or "-",
                "gender": u.gender or "P",
                "prodi": u.department or "Sistem Informasi Geografis",
                "institution": u.institution or "Sekolah Vokasi - UGM",
                "email": u.email,
                "username": u.username,
                "grids": grid_items
            })

    selected_area_name = study_area_map.get(study_area_id, "Seluruh Wilayah") if study_area_id else "Seluruh Wilayah"

    return {
        "status": "success",
        "selected_area_name": selected_area_name,
        "selected_year": year,
        "summary": {
            "total_mappers": total_mappers_count,
            "total_grids": total_grids_count,
            "in_progress_count": in_progress_count,
            "review_qc_count": review_qc_count,
            "qc1_count": qc1_count,
            "qc2_count": qc2_count,
            "finishing_count": finishing_count,
            "completion_percentage": round((finishing_count / total_grids_count * 100), 1) if total_grids_count > 0 else 0
        },
        "study_areas": [{"id": sa.id, "name": sa.name} for sa in study_areas],
        "years": distinct_years,
        "data": records
    }

@router.patch("/tasks/{task_id}/stages")
def update_task_stage(
    task_id: int,
    payload: TaskStageUpdateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_reviewer)
):
    """
    Memperbarui status tahapan QC (QC 1, QC 2, Finishing, Keterangan)
    oleh Dosen, Supervisi, atau Lead Admin.
    """
    task = db.query(TaskGrid).filter(TaskGrid.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task Grid tidak ditemukan.")

    if payload.qc1_approved is not None:
        task.qc1_approved = payload.qc1_approved
        if payload.qc1_approved and task.status != TaskStatus.APPROVED.value:
            task.status = TaskStatus.APPROVED.value
    
    if payload.qc2_approved is not None:
        task.qc2_approved = payload.qc2_approved

    if payload.finishing_approved is not None:
        task.finishing_approved = payload.finishing_approved

    if payload.reviewer_notes is not None:
        task.reviewer_notes = payload.reviewer_notes

    if payload.status is not None:
        task.status = payload.status

    task.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(task)

    return {
        "status": "success",
        "message": f"Tahapan QC untuk {task.grid_code} berhasil diperbarui.",
        "data": {
            "task_id": task.id,
            "grid_code": task.grid_code,
            "status": task.status,
            "qc1_approved": task.qc1_approved,
            "qc2_approved": task.qc2_approved,
            "finishing_approved": task.finishing_approved,
            "reviewer_notes": task.reviewer_notes
        }
    }

@router.get("/export/excel")
def export_progress_excel(
    study_area_id: Optional[int] = None,
    year: Optional[int] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Menghasilkan file Excel (.xlsx) dengan tata letak & merge cell identik
    dengan template Google Sheets monitoring annotasi.
    """
    data_res = get_progress_table(study_area_id=study_area_id, year=year, search=search, db=db, current_user=current_user)
    records = data_res["data"]
    selected_area = data_res.get("selected_area_name", "Seluruh Wilayah")
    if not HAS_OPENPYXL:
        raise HTTPException(
            status_code=501,
            detail="Fitur ekspor Excel (.xlsx) memerlukan pustaka openpyxl di server. Silakan gunakan format CSV atau hubungi administrator."
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Progres Training Sample"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    sub_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    area_header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")

    white_font_bold = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="0F172A")
    true_font = Font(name="Calibri", size=10, bold=True, color="166534")
    false_font = Font(name="Calibri", size=10, color="94A3B8")

    true_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    false_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Header Row 1 & 2
    ws.merge_cells("A1:A2")
    ws["A1"] = "Nama"
    ws.merge_cells("B1:B2")
    ws["B1"] = "NIM"
    ws.merge_cells("C1:C2")
    ws["C1"] = "L/P"
    ws.merge_cells("D1:D2")
    ws["D1"] = "Program Studi"

    ws.merge_cells("E1:F1")
    ws["E1"] = selected_area if selected_area != "Seluruh Wilayah" else "Wilayah Spasial"
    ws["E2"] = "GRID"
    ws["F2"] = "Tahun"

    ws.merge_cells("G1:G2")
    ws["G1"] = "Status Dalam Proses"
    ws.merge_cells("H1:H2")
    ws["H1"] = "Status Review QC"
    ws.merge_cells("I1:I2")
    ws["I1"] = "QC 1 (Approved)"
    ws.merge_cells("J1:J2")
    ws["J1"] = "QC 2 (Approved)"
    ws.merge_cells("K1:K2")
    ws["K1"] = "Finishing\n(Mas Danang & Mas Habib)"
    ws.merge_cells("L1:L2")
    ws["L1"] = "Keterangan"

    for row in range(1, 3):
        for col in range(1, 13):
            cell = ws.cell(row=row, column=col)
            cell.font = white_font_bold
            cell.border = cell_border
            cell.alignment = align_center
            if col in [5, 6] and row == 1:
                cell.fill = area_header_fill
            elif row == 2 and col in [5, 6]:
                cell.fill = sub_header_fill
            else:
                cell.fill = header_fill

    current_row = 3

    for rec in records:
        grids = rec.get("grids", [])
        if not grids:
            continue

        start_row = current_row
        end_row = current_row + len(grids) - 1

        for g_idx, g in enumerate(grids):
            row_num = current_row + g_idx

            ws.cell(row=row_num, column=1, value=rec["nama"])
            ws.cell(row=row_num, column=2, value=rec["nim"])
            ws.cell(row=row_num, column=3, value=rec["gender"])
            ws.cell(row=row_num, column=4, value=rec["prodi"])

            c_grid = ws.cell(row=row_num, column=5, value=g["grid_number"] or g["grid_code"])
            c_year = ws.cell(row=row_num, column=6, value=g["year"])
            
            c_in_progress = ws.cell(row=row_num, column=7, value="TRUE" if g["in_progress"] else "FALSE")
            c_review_qc = ws.cell(row=row_num, column=8, value="TRUE" if g["review_qc"] else "FALSE")
            c_qc1 = ws.cell(row=row_num, column=9, value="TRUE" if g["qc1_approved"] else "FALSE")
            c_qc2 = ws.cell(row=row_num, column=10, value="TRUE" if g["qc2_approved"] else "FALSE")
            c_finishing = ws.cell(row=row_num, column=11, value="TRUE" if g["finishing_approved"] else "FALSE")
            c_notes = ws.cell(row=row_num, column=12, value=g["keterangan"])

            for col in range(1, 13):
                c = ws.cell(row=row_num, column=col)
                c.border = cell_border
                c.font = data_font

            ws.cell(row=row_num, column=1).alignment = align_left
            ws.cell(row=row_num, column=2).alignment = align_center
            ws.cell(row=row_num, column=3).alignment = align_center
            ws.cell(row=row_num, column=4).alignment = align_left
            c_grid.alignment = align_center
            c_year.alignment = align_center
            c_notes.alignment = align_left

            for bool_cell in [c_in_progress, c_review_qc, c_qc1, c_qc2, c_finishing]:
                bool_cell.alignment = align_center
                if bool_cell.value == "TRUE":
                    bool_cell.font = true_font
                    bool_cell.fill = true_fill
                else:
                    bool_cell.font = false_font
                    bool_cell.fill = false_fill

        if end_row > start_row:
            ws.merge_cells(start_row=start_row, end_row=end_row, start_column=1, end_column=1)
            ws.merge_cells(start_row=start_row, end_row=end_row, start_column=2, end_column=2)
            ws.merge_cells(start_row=start_row, end_row=end_row, start_column=3, end_column=3)
            ws.merge_cells(start_row=start_row, end_row=end_row, start_column=4, end_column=4)

        current_row = end_row + 1

    column_widths = {
        "A": 26, "B": 24, "C": 6, "D": 32, "E": 10, "F": 8,
        "G": 18, "H": 16, "I": 16, "J": 16, "K": 26, "L": 28
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 20

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Rekap_Progres_Training_Sample_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/export/csv")
def export_progress_csv(
    study_area_id: Optional[int] = None,
    year: Optional[int] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Ekspor data rekapitulasi progres monitoring dalam format CSV standar.
    """
    data_res = get_progress_table(study_area_id=study_area_id, year=year, search=search, db=db, current_user=current_user)
    records = data_res["data"]
    selected_area = data_res["selected_area_name"]

    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "Nama",
        "NIM",
        "L/P",
        "Program Studi",
        f"{selected_area} - GRID",
        "Tahun",
        "Status Dalam Proses",
        "Status Review QC",
        "QC 1 (Approved)",
        "QC 2 (Approved)",
        "Finishing (Mas Danang & Mas Habib)",
        "Keterangan"
    ])

    for rec in records:
        for g in rec.get("grids", []):
            writer.writerow([
                rec["nama"],
                rec["nim"],
                rec["gender"],
                rec["prodi"],
                g["grid_number"] or g["grid_code"],
                g["year"],
                "TRUE" if g["in_progress"] else "FALSE",
                "TRUE" if g["review_qc"] else "FALSE",
                "TRUE" if g["qc1_approved"] else "FALSE",
                "TRUE" if g["qc2_approved"] else "FALSE",
                "TRUE" if g["finishing_approved"] else "FALSE",
                g["keterangan"]
            ])

    output.seek(0)
    filename = f"Rekap_Progres_Training_Sample_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
